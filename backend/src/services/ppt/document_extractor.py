"""Extract text content from PDF, DOCX, TXT, MD files.

Inspired by ppt-master's source_to_md pipeline.
Uses PyMuPDF for PDF, python-docx for DOCX, plain read for text.
"""

from __future__ import annotations

import re
from pathlib import Path
from dataclasses import dataclass, field


@dataclass
class ExtractedSection:
    """A section extracted from a document."""
    heading: str = ""
    level: int = 0  # 0=body, 1=h1, 2=h2, 3=h3
    text: str = ""
    page_num: int = 0


@dataclass
class ExtractedDocument:
    """Full extracted document content."""
    title: str = ""
    sections: list[ExtractedSection] = field(default_factory=list)
    total_pages: int = 0
    source_format: str = ""

    @property
    def full_text(self) -> str:
        parts = []
        for s in self.sections:
            if s.heading:
                prefix = "#" * s.level if s.level else ""
                parts.append(f"{prefix} {s.heading}" if prefix else s.heading)
            if s.text:
                parts.append(s.text)
        return "\n\n".join(parts)

    def to_markdown(self) -> str:
        """Convert to structured markdown for LLM consumption."""
        parts = []
        if self.title:
            parts.append(f"# {self.title}\n")
        for s in self.sections:
            if s.heading:
                prefix = "#" * max(s.level, 1)
                parts.append(f"{prefix} {s.heading}")
            if s.text:
                parts.append(s.text.strip())
        return "\n\n".join(parts)


def extract_pdf(file_path: str) -> ExtractedDocument:
    """Extract text from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("PyMuPDF required: pip install PyMuPDF")

    doc = fitz.open(file_path)
    result = ExtractedDocument(source_format="pdf", total_pages=len(doc))

    # Analyze font sizes for heading detection
    from collections import Counter
    size_counter: Counter = Counter()

    for page in doc:
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            if block.get("type") == 0:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        size = round(span.get("size", 12), 1)
                        text = span.get("text", "").strip()
                        if text:
                            size_counter[size] += len(text)

    if not size_counter:
        doc.close()
        return result

    sorted_sizes = sorted(size_counter.items(), key=lambda x: x[1], reverse=True)
    body_size = sorted_sizes[0][0]
    larger_sizes = sorted([s for s, _ in sorted_sizes if s > body_size + 1], reverse=True)

    h1_size = larger_sizes[0] if len(larger_sizes) >= 1 else body_size + 6
    h2_size = larger_sizes[1] if len(larger_sizes) >= 2 else body_size + 3
    h3_size = larger_sizes[2] if len(larger_sizes) >= 3 else body_size + 1.5

    def get_heading_level(size: float, text: str) -> int:
        if size >= h1_size and len(text) < 100:
            return 1
        if size >= h2_size and len(text) < 120:
            return 2
        if size >= h3_size and len(text) < 150:
            return 3
        return 0

    current_section = ExtractedSection()

    for page_idx, page in enumerate(doc):
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                line_text = ""
                max_size = 0
                for span in line.get("spans", []):
                    line_text += span.get("text", "")
                    max_size = max(max_size, span.get("size", 12))

                line_text = line_text.strip()
                if not line_text:
                    continue

                level = get_heading_level(max_size, line_text)

                if level > 0:
                    # Save previous section
                    if current_section.text.strip() or current_section.heading:
                        result.sections.append(current_section)
                    current_section = ExtractedSection(
                        heading=line_text, level=level, page_num=page_idx + 1
                    )
                else:
                    if current_section.text:
                        current_section.text += " " + line_text
                    else:
                        current_section.text = line_text

    if current_section.text.strip() or current_section.heading:
        result.sections.append(current_section)

    # Set title from first h1
    for s in result.sections:
        if s.level == 1:
            result.title = s.heading
            break

    doc.close()
    return result


def extract_docx(file_path: str) -> ExtractedDocument:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError("python-docx required: pip install python-docx")

    doc = DocxDocument(file_path)
    result = ExtractedDocument(source_format="docx")

    current_section = ExtractedSection()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name.lower() if para.style else ""

        # Detect headings by style
        level = 0
        if "heading 1" in style_name or "title" in style_name:
            level = 1
        elif "heading 2" in style_name:
            level = 2
        elif "heading 3" in style_name:
            level = 3

        if level > 0:
            if current_section.text.strip() or current_section.heading:
                result.sections.append(current_section)
            current_section = ExtractedSection(heading=text, level=level)
            if level == 1 and not result.title:
                result.title = text
        else:
            if current_section.text:
                current_section.text += " " + text
            else:
                current_section.text = text

    if current_section.text.strip() or current_section.heading:
        result.sections.append(current_section)

    return result


def extract_text(file_path: str) -> ExtractedDocument:
    """Extract from plain text / markdown files."""
    path = Path(file_path)
    content = path.read_text(encoding="utf-8", errors="replace")
    result = ExtractedDocument(source_format=path.suffix.lstrip("."))

    # Try to detect markdown headings
    lines = content.split("\n")
    current_section = ExtractedSection()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Markdown heading detection
        heading_match = re.match(r'^(#{1,3})\s+(.+)', stripped)
        if heading_match:
            if current_section.text.strip() or current_section.heading:
                result.sections.append(current_section)
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            current_section = ExtractedSection(heading=heading_text, level=level)
            if level == 1 and not result.title:
                result.title = heading_text
        else:
            if current_section.text:
                current_section.text += " " + stripped
            else:
                current_section.text = stripped

    if current_section.text.strip() or current_section.heading:
        result.sections.append(current_section)

    if not result.title and result.sections:
        result.title = result.sections[0].heading or "Untitled"

    return result


def extract_document(file_path: str) -> ExtractedDocument:
    """Auto-detect format and extract text."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    extractors = {
        ".pdf": extract_pdf,
        ".docx": extract_docx,
        ".doc": extract_docx,
    }

    extractor = extractors.get(suffix, extract_text)
    return extractor(file_path)
