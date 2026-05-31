"""Excel Builder — generate Excel files from templates and data.

Uses openpyxl to create real editable XLSX with formatting, formulas, charts.
"""

from __future__ import annotations

import io
from pathlib import Path
from dataclasses import dataclass, field
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, PieChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl required: pip install openpyxl")


@dataclass
class ColumnDef:
    name: str
    width: int = 15
    header_fill: str = "2B579A"
    header_font_color: str = "FFFFFF"
    number_format: str = "General"


@dataclass
class SheetDef:
    name: str = "Sheet1"
    columns: list[ColumnDef] = field(default_factory=list)
    data: list[list[Any]] = field(default_factory=list)
    chart_type: str | None = None  # "bar", "pie", "line"
    chart_title: str = ""
    chart_data_col: int = 1  # column index for chart values


@dataclass
class ExcelPlan:
    title: str = "Report"
    sheets: list[SheetDef] = field(default_factory=list)
    author: str = "VietRAG"


def _apply_header_style(cell, fill_color: str, font_color: str):
    cell.font = Font(bold=True, color=font_color, size=11)
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        bottom=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
    )


def _add_chart(ws, chart_type: str, title: str, data_col: int, num_rows: int, num_cols: int):
    chart_map = {"bar": BarChart, "pie": PieChart, "line": LineChart}
    ChartClass = chart_map.get(chart_type, BarChart)
    chart = ChartClass()
    chart.title = title
    chart.style = 10
    chart.width = 20
    chart.height = 12

    cats = Reference(ws, min_col=1, min_row=2, max_row=num_rows)
    vals = Reference(ws, min_col=data_col + 1, min_row=1, max_row=num_rows)
    chart.add_data(vals, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{num_rows + 3}")


def build_excel(plan: ExcelPlan, output_path: str | None = None) -> bytes:
    """Build an Excel file from an ExcelPlan. Returns file bytes."""
    wb = Workbook()
    wb.properties.creator = plan.author
    wb.properties.title = plan.title

    for i, sheet_def in enumerate(plan.sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sheet_def.name

        # Headers
        for col_idx, col_def in enumerate(sheet_def.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_def.name)
            _apply_header_style(cell, col_def.header_fill, col_def.header_font_color)
            ws.column_dimensions[get_column_letter(col_idx)].width = col_def.width

        # Data rows
        for row_idx, row_data in enumerate(sheet_def.data, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if col_idx <= len(sheet_def.columns):
                    cell.number_format = sheet_def.columns[col_idx - 1].number_format
                # Alternate row coloring
                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

        # Chart
        if sheet_def.chart_type and sheet_def.data:
            _add_chart(ws, sheet_def.chart_type, sheet_def.chart_title,
                      sheet_def.chart_data_col, len(sheet_def.data) + 1, len(sheet_def.columns))

        # Auto-filter
        if sheet_def.columns and sheet_def.data:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet_def.columns))}{len(sheet_def.data) + 1}"

        # Freeze header row
        ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    xlsx_bytes = buffer.read()

    if output_path:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        Path(output_path).write_bytes(xlsx_bytes)

    return xlsx_bytes


# === Predefined templates ===

TEMPLATES = {
    "invoice": {
        "title": "Invoice",
        "sheets": [{
            "name": "Invoice",
            "columns": [
                ColumnDef("STT", 8), ColumnDef("San pham", 30),
                ColumnDef("So luong", 12, number_format="#,##0"),
                ColumnDef("Don gia", 15, number_format="#,##0"),
                ColumnDef("Thanh tien", 18, number_format="#,##0"),
            ],
            "data": [
                [1, "San pham A", 10, 100000, "=C2*D2"],
                [2, "San pham B", 5, 200000, "=C3*D3"],
                [3, "San pham C", 3, 150000, "=C4*D4"],
                ["", "", "", "Tong cong:", "=SUM(E2:E4)"],
            ],
        }],
    },
    "budget": {
        "title": "Budget Report",
        "sheets": [{
            "name": "Budget",
            "columns": [
                ColumnDef("Muc", 25), ColumnDef("Quy 1", 15, number_format="#,##0"),
                ColumnDef("Quy 2", 15, number_format="#,##0"),
                ColumnDef("Quy 3", 15, number_format="#,##0"),
                ColumnDef("Quy 4", 15, number_format="#,##0"),
                ColumnDef("Tong", 18, number_format="#,##0"),
            ],
            "data": [
                ["Luong nhan vien", 50000000, 50000000, 55000000, 55000000, "=SUM(B2:E2)"],
                ["Marketing", 10000000, 15000000, 12000000, 20000000, "=SUM(B3:E3)"],
                ["Van phong", 5000000, 5000000, 5000000, 5000000, "=SUM(B4:E4)"],
                ["Cong nghe", 8000000, 8000000, 10000000, 10000000, "=SUM(B5:E5)"],
                ["Tong chi phi", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", "=SUM(E2:E5)", "=SUM(F2:F5)"],
            ],
            "chart_type": "bar",
            "chart_title": "Chi phi theo quy",
        }],
    },
    "employee": {
        "title": "Danh sach nhan vien",
        "sheets": [{
            "name": "Nhan vien",
            "columns": [
                ColumnDef("STT", 8), ColumnDef("Ho ten", 25),
                ColumnDef("Phong ban", 20), ColumnDef("Chuc vu", 20),
                ColumnDef("Luong", 18, number_format="#,##0"),
                ColumnDef("Ngay vao lam", 15, number_format="DD/MM/YYYY"),
            ],
            "data": [
                [1, "Nguyen Van A", "IT", "Developer", 15000000, "01/01/2024"],
                [2, "Tran Thi B", "Marketing", "Manager", 20000000, "15/03/2023"],
                [3, "Le Van C", "Sales", "Staff", 12000000, "20/06/2024"],
            ],
        }],
    },
    "inventory": {
        "title": "Quan ly kho",
        "sheets": [{
            "name": "Ton kho",
            "columns": [
                ColumnDef("Ma SP", 12), ColumnDef("Ten san pham", 30),
                ColumnDef("Ton dau", 12, number_format="#,##0"),
                ColumnDef("Nhap", 12, number_format="#,##0"),
                ColumnDef("Xuat", 12, number_format="#,##0"),
                ColumnDef("Ton cuoi", 12, number_format="#,##0"),
                ColumnDef("Don gia", 15, number_format="#,##0"),
                ColumnDef("Thanh tien", 18, number_format="#,##0"),
            ],
            "data": [
                ["SP001", "Laptop Dell", 50, 100, 80, "=C2+D2-E2", 15000000, "=F2*G2"],
                ["SP002", "Chuot Logitech", 200, 500, 350, "=C3+D3-E3", 300000, "=F3*G3"],
                ["SP003", "Ban phim", 100, 200, 150, "=C4+D4-E4", 500000, "=F4*G4"],
            ],
            "chart_type": "bar",
            "chart_title": "Ton kho san pham",
        }],
    },
}


def build_from_template(template_name: str, output_path: str | None = None) -> bytes:
    """Build Excel from a predefined template."""
    tmpl = TEMPLATES.get(template_name)
    if not tmpl:
        raise ValueError(f"Template not found: {template_name}. Available: {list(TEMPLATES.keys())}")

    plan = ExcelPlan(
        title=tmpl["title"],
        sheets=[SheetDef(**s) for s in tmpl["sheets"]],
    )
    return build_excel(plan, output_path)
